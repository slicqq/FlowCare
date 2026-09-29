"""Add the FlowCare logo and real product screenshots to the pitch deck.

- Keys the white background out of the logo master so the mark sits on any fill.
- Large mark on the cover and the closing slide; small mark on every other slide.
- Fills the slide-5 "[Paste screenshot]" placeholder with two real screenshots.
- Inserts two new screenshot slides cloned from the template chrome.
- Renumbers the page numbers afterwards.

Screenshots are captured from the running app by shots-deck.py. Nothing here is
a mock-up or a redraw.
"""
import copy
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

DECK = "/home/user/deck"
SRC = f"{DECK}/FlowCare-CuriousParc-2026.pptx"
OUT = f"{DECK}/FlowCare-CuriousParc-2026.pptx"
SHOTS = f"{DECK}/screens"
BUILD = f"{DECK}/.build"
os.makedirs(BUILD, exist_ok=True)

BRAND = (0x3D, 0xBB, 0xB9)
INK = RGBColor(0x18, 0x27, 0x44)
GREY = RGBColor(0x65, 0x73, 0x85)
HAIRLINE = RGBColor(0xD8, 0xDF, 0xE7)
FONT = "Aptos"
FONT_B = "Aptos Bold"


# ---------------------------------------------------------------- logo asset
def make_transparent_mark(src, dst):
    """Un-composite a solid-colour mark from its white background.

    The master is teal on opaque white. A pixel is P = a*C + (1-a)*255 per
    channel, so alpha recovers exactly from the channel with the most contrast
    against white (red: 61 vs 255). Thresholding instead would leave a white
    fringe on the antialiased curve.
    """
    im = Image.open(src).convert("RGB")
    w, h = im.size
    out = Image.new("RGBA", (w, h))
    px, op = im.load(), out.load()
    cr = BRAND[0]
    for y in range(h):
        for x in range(w):
            r, g, b = px[x, y]
            a = (255 - r) / (255 - cr)
            # Clamp off sensor/encoder noise, otherwise a stray 254-grey pixel
            # keeps the alpha bounding box at the full canvas and the trim is
            # a no-op, leaving the mark padded and off-ratio.
            a = 0 if a < 0.04 else (1 if a > 1 else a)
            op[x, y] = (*BRAND, int(round(a * 255)))
    bbox = out.getchannel("A").getbbox()
    out = out.crop(bbox)
    out.save(dst)
    return out.size


MARK = f"{BUILD}/mark.png"
MW, MH = make_transparent_mark("/home/user/flowcare/public/brand/logo-source.png", MARK)
ASPECT = MH / MW
print(f"mark trimmed to {MW}x{MH}  (h = w * {ASPECT:.4f})")


def prep_shot(name, target_w=1400):
    """Downscale a 2x screenshot so the deck does not balloon."""
    src = f"{SHOTS}/{name}"
    dst = f"{BUILD}/{name}"
    im = Image.open(src).convert("RGB")
    if im.width > target_w:
        im = im.resize((target_w, round(im.height * target_w / im.width)), Image.LANCZOS)
    im.save(dst, optimize=True)
    return dst, im.width / im.height


# ------------------------------------------------------------------- helpers
def set_text(shape, text):
    """Replace a textbox's text, keeping the first run's formatting."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        p0.add_run()
    p0.runs[0].text = text
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def add_mark(slide, left, width):
    pic = slide.shapes.add_picture(MARK, Inches(left), Inches(0), width=Inches(width))
    pic.height = Emu(int(pic.width * ASPECT))
    return pic


def add_caption(slide, left, top, width, label, sub=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.5))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = label
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.name = FONT_B
    r.font.color.rgb = INK
    if sub:
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = sub
        r2.font.size = Pt(10.5)
        r2.font.name = FONT
        r2.font.color.rgb = GREY
    return box


def add_shot(slide, name, left, top, box_w, box_h):
    """Fit a screenshot inside a box, centred, with a hairline border."""
    path, ar = prep_shot(name)
    w, h = box_w, box_w / ar
    if h > box_h:
        h, w = box_h, box_h * ar
    x = left + (box_w - w) / 2
    y = top + (box_h - h) / 2
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    pic.line.color.rgb = HAIRLINE
    pic.line.width = Pt(0.75)
    return pic


def clone_slide(prs, src_slide):
    """Duplicate a slide's shape tree onto a new Blank slide, remapping rels."""
    blank = next(l for l in prs.slide_layouts if l.name == "Blank")
    new = prs.slides.add_slide(blank)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    for sh in src_slide.shapes:
        new.shapes._spTree.append(copy.deepcopy(sh._element))
    for el in new.shapes._spTree.iter():
        for attr in (qn("r:embed"), qn("r:link")):
            rid = el.get(attr)
            if rid:
                rel = src_slide.part.rels[rid]
                el.set(attr, new.part.relate_to(rel.target_part, rel.reltype))
    return new


def strip_to_chrome(slide, keep_ids):
    for sh in list(slide.shapes):
        if sh.shape_id not in keep_ids:
            sh._element.getparent().remove(sh._element)


def move_slide(prs, frm, to):
    lst = prs.slides._sldIdLst
    el = list(lst)[frm]
    lst.remove(el)
    lst.insert(to, el)


# ---------------------------------------------------------------------- build
prs = Presentation(SRC)
slides = list(prs.slides)
print(f"loaded {len(slides)} slides")

# 1. Cover: large mark in the open lower-right quadrant.
cover = slides[0]
m = add_mark(cover, 13.75, 3.40)
m.top = Inches(6.15)
print(f"cover mark  {m.width/914400:.2f} x {m.height/914400:.2f} in")

# 2. Closing slide: the contact block runs to x=9.17 and the partner card
# starts at x=12.34, so the mark goes in the gap between them.
closing = slides[-1]
m = add_mark(closing, 9.50, 2.20)
m.top = Inches(6.15)
print(f"closing mark {m.width/914400:.2f} x {m.height/914400:.2f} in")

# 3. Small mark on the content slides, left of the CP badge at 17.18in.
for s in slides[1:-1]:
    m = add_mark(s, 16.28, 0.66)
    m.top = Inches(0.60)

# 4. Slide 5 — replace the paste-a-screenshot placeholder with two real ones.
s5 = slides[4]
by_id = {sh.shape_id: sh for sh in s5.shapes}
set_text(by_id[35], "PROTOTYPE — RUNS LOCALLY, TESTED, NOT YET DEPLOYED")
by_id[35].top = Inches(4.32)
by_id[36]._element.getparent().remove(by_id[36]._element)  # placeholder text
add_shot(s5, "06-patient-dash.png", 1.50, 4.95, 5.00, 3.13)
add_shot(s5, "08-hospital-detail.png", 6.86, 4.95, 5.00, 3.13)
set_text(
    by_id[37],
    "Patient dashboard and a hospital profile, captured from the running app. "
    "The banner across the top is the app's own warning that this instance is "
    "serving synthetic records, not real hospitals or real patients.",
)
by_id[37].top = Inches(8.28)
by_id[37].width = Inches(9.94)

# 5. Two new screenshot slides, cloned from slide 5's chrome.
CHROME = {2, 4, 5, 6, 8, 11, 13, 54}

ROW_X = [1.07, 5.62, 10.17, 14.72]
CELL_W = 4.19

PATIENT = [
    ("01-landing.png", "Landing", "The pitch, the journey, and who it is for"),
    ("02-role-choice.png", "Choosing a role", "Patient or hospital staff — separate front doors"),
    ("03-patient-signup.png", "Patient sign-up", "Never asks a patient for hospital admin fields"),
    ("06-patient-dash.png", "Patient dashboard", "Current appointment, queue, recent activity"),
    ("07-discovery.png", "Discovery", "Search and filters that can actually be enforced"),
    ("08-hospital-detail.png", "Hospital profile", "Every fact dated and attributed to its source"),
    ("09-booking.png", "Requesting a slot", "A request, not a confirmed booking"),
    ("11-assistant.png", "Assistant", "Proposes filters; it never diagnoses or prescribes"),
]

STAFF = [
    ("13-staff-dash.png", "Staff dashboard", "Today's load and what is waiting on the hospital"),
    ("14-staff-queue.png", "Queue management", "Check in, call through, close off"),
    ("15-staff-appts.png", "Appointments", "Confirm or decline the requests patients sent"),
    ("16-staff-team.png", "Staff access", "Admins approve access; nothing is auto-granted"),
]

MOBILE = [
    ("17-landing-mobile.png", "Landing"),
    ("19-role-mobile.png", "Role choice"),
    ("18-patient-mobile.png", "Dashboard"),
]


def new_screens_slide(title, subtitle):
    s = clone_slide(prs, s5)
    strip_to_chrome(s, CHROME)
    ids = {sh.shape_id: sh for sh in s.shapes}
    set_text(ids[4], title)
    set_text(ids[5], subtitle)
    add_mark(s, 16.28, 0.66).top = Inches(0.60)
    return s


# -- patient slide: two rows of four
sa = new_screens_slide(
    "The Product — What a Patient Sees",
    "Screenshots from the running app. The yellow bar is the app labelling its own demo data.",
)
for i, (name, label, sub) in enumerate(PATIENT):
    col, row = i % 4, i // 4
    x = ROW_X[col]
    y = 2.35 + row * 3.85
    add_shot(sa, name, x, y, CELL_W, 2.62)
    add_caption(sa, x, y + 2.70, CELL_W, label, sub)

# -- staff slide: one row of four, then the phone layouts
# Title cap for this box is ~46 characters at 37.5pt bold in a 13.02in frame;
# anything longer wraps onto the subtitle. The phone story moves to the sub.
sb = new_screens_slide(
    "The Product — The Hospital Side",
    "Staff screens are permission-checked on the server. The same design collapses to a bottom nav on a phone.",
)
for i, (name, label, sub) in enumerate(STAFF):
    x = ROW_X[i]
    add_shot(sb, name, x, 2.35, CELL_W, 2.62)
    add_caption(sb, x, 5.05, CELL_W, label, sub)

MOB_W, MOB_H = 1.95, 4.05
for i, (name, label) in enumerate(MOBILE):
    x = 1.07 + i * (MOB_W + 0.42)
    add_shot(sb, name, x, 6.05, MOB_W, MOB_H)
    add_caption(sb, x, 10.18, MOB_W, label)

note = sb.shapes.add_textbox(Inches(8.60), Inches(6.30), Inches(10.33), Inches(3.0))
tf = note.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_right = 0
lines = [
    ("One design language, three widths.",
     "Desktop keeps the full navigation, tablet drops to a condensed header, and "
     "phones move the primary navigation to a bottom bar within thumb reach."),
    ("Nothing is hidden to enforce a rule.",
     "The staff screens check the session, the role and the hospital on the server "
     "on every request. Hiding a button is presentation, not authorisation."),
    ("Loading states have a shape.",
     "Each screen has a skeleton built to the proportions of its real content, so "
     "the page does not jump when the data lands."),
]
first = True
for head, body in lines:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(9)
    r = p.add_run()
    r.text = head + "  "
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.name = FONT_B
    r.font.color.rgb = INK
    r2 = p.add_run()
    r2.text = body
    r2.font.size = Pt(13)
    r2.font.name = FONT
    r2.font.color.rgb = GREY

# order: put the two new slides straight after slide 5
n = len(prs.slides._sldIdLst)
move_slide(prs, n - 2, 5)  # sa
move_slide(prs, n - 1, 6)  # sb

# 6. Renumber pages.
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and abs(sh.left - Inches(18.11)) < Inches(0.12) \
                and abs(sh.top - Inches(10.73)) < Inches(0.12):
            set_text(sh, f"{i:02d}")

prs.save(OUT)
print(f"saved {OUT}  ({len(prs.slides._sldIdLst)} slides, {os.path.getsize(OUT)/1e6:.1f} MB)")
